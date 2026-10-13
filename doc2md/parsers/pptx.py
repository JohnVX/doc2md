"""PowerPoint(pptx)解析: python-pptx, 递归分组图形 + 内嵌图提取+OCR占位.

通用:
  - 递归 walk 进 GROUP, 不丢分组内内容
  - 标题: slide.shapes.title -> 否则首个非空文本框
  - 文本框/表格/图片/备注; 跳过母版占位符样板文本
  - 图片 blob -> 存 assets -> OCR 占位(收尾执行)
"""
from pathlib import Path

from .. import util
from ._common import defer, is_placeholder, md_table


def parse(path, ctx=None):
    """pptx 解析入口(分组递归+图片OCR)."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(path)
    stem = Path(path).stem

    core_title = None
    try:
        core_title = prs.core_properties.title
    except Exception:
        pass
    first_slide_title = ""
    if prs.slides:
        try:
            first_slide_title = _slide_title(prs.slides[0])[0]
        except Exception:
            first_slide_title = ""
    title = util.pick_title([first_slide_title, core_title], stem)

    out = []
    deferred = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(_render_slide(i, slide, ctx, deferred))
    body = "\n\n".join(x for x in out if x) or f"<!-- 空 presentation: {Path(path).name} -->"
    return {"title": title, "body": body, "tags": [], "meta": {}, "deferred": deferred}


def _walk(shapes):
    """递归展开分组图形. 用 hasattr(sh,'shapes') 判 GroupShape, 不依赖
    模块级 MSO_SHAPE_TYPE(其在 parse() 内才 import, 否则 NameError 被吞致分组不递归)."""
    for sh in shapes:
        try:
            if hasattr(sh, "shapes"):
                yield from _walk(sh.shapes)
            else:
                yield sh
        except Exception:
            yield sh


def _slide_title(slide):
    """取幻灯片标题. 返回 (标题文本, 来源 shape 或 None).

    来源 shape 供 _render_slide 跳过, 避免兜底取用的首个文本框又在正文
    重复输出一次.
    """
    try:
        t = slide.shapes.title
        if t is not None and t.text.strip():
            return t.text.strip(), t
    except Exception:
        pass
    # 兜底: 首个非空文本框, 仅当其为**单行**时视为标题 (多行文本框是正文,
    # 不能拿首行当标题而丢掉其余行). 返回来源 shape 供跳过.
    for sh in _walk(slide.shapes):
        try:
            if sh.has_text_frame and sh.text_frame.text.strip():
                lines = [x for x in sh.text_frame.text.splitlines() if x.strip()]
                if len(lines) == 1:
                    return lines[0].strip(), sh
        except Exception:
            continue
    return "", None


def _same_shape(a, b):
    """两个 shape 是否指同一对象.

    python-pptx 每次遍历/属性访问都返回**新包装对象**, `a is b` 永远为假;
    必须比较底层 XML element(_element 稳定), 否则标题 placeholder 会被重复输出.
    """
    if a is None or b is None:
        return False
    ea = getattr(a, "_element", None)
    eb = getattr(b, "_element", None)
    if ea is not None and eb is not None:
        return ea is eb
    return a is b


def _render_slide(idx, slide, ctx, deferred):
    """渲染一张幻灯片为 md."""
    title, title_shape = _slide_title(slide)
    head = f"## Slide {idx}" + (f": {title}" if title else "")

    items = []
    img_no = 0
    unhandled = 0
    for sh in _walk(slide.shapes):
        if _same_shape(sh, title_shape):
            continue
        handled = False
        if sh.has_text_frame:
            txt = _fmt_textframe(sh.text_frame)
            if txt and not is_placeholder(txt):
                items.append(txt)
            handled = True
        if sh.has_table:
            items.append(_render_table(sh.table))
            handled = True
        try:
            if sh.shape_type == 13:  # PICTURE
                img_no += 1
                _handle_picture(sh, img_no, ctx, items, deferred)
                handled = True
        except Exception:
            pass
        if not handled:
            # 排除纯装饰形状: LINE(连接符/箭头)/FREEFORM(自由图形) 无内容载体,
            # 计入 chart defer 会让 stage2 误以为有大量图表待处理 (实际只是箭头).
            # 保留 CHART/EMBEDDED_OLE_OBJECT 等真内容载体.
            # 注意: MSO_SHAPE_TYPE 必须在此局部 import (模块级 import 在 parse() 内,
            # 否则此处 NameError 被 except 吞掉 —— 同 _walk 的坑).
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            st = None
            try:
                st = sh.shape_type
            except Exception:
                pass
            if st not in (MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.FREEFORM):
                unhandled += 1
    if unhandled:
        items.append(defer(deferred, "chart", unhandled, "图表/SmartArt等未提取"))

    # 备注
    try:
        nt = slide.notes_slide.notes_text_frame.text.strip()
        if nt and not is_placeholder(nt):
            items.append(f"> 备注: {nt}")
    except Exception:
        pass

    return head + "\n\n" + "\n\n".join(items) if items else head


def _fmt_textframe(tf):
    """格式化文本框为 md 段落."""
    out = []
    for para in tf.paragraphs:
        t = "".join(r.text or "" for r in para.runs) or para.text
        t = t.strip()
        if not t:
            continue
        lvl = getattr(para, "level", 0) or 0
        if lvl > 0:
            out.append("  " * lvl + "- " + t)
        else:
            out.append(t)
    return "\n".join(out)


def _render_table(tbl):
    """渲染 pptx 表格为 md 表格."""
    try:
        ncols = len(tbl.columns)
    except Exception:
        ncols = max((len(r.cells) for r in tbl.rows), default=0)
    if ncols < 1 or not tbl.rows:
        return ""
    rows = []
    for r in tbl.rows:
        # 不在此转义: 提取原始文本, 转义统一由 md_table 的 esc 处理
        # (避免 | 被双重转义 => \\\|)
        cells = [c.text for c in r.cells]
        rows.append(cells)
    return md_table(rows, ncols)


def _handle_picture(sh, no, ctx, items, deferred):
    """处理图片形状: 提取 blob 存 assets, 插入 OCR 占位(OCR 在 pipeline 收尾执行).

    worker 进程不再加载 OCR 引擎, 避免多进程内存爆炸 (每进程 ~400MB).
    """
    try:
        blob = sh.image.blob
        ext = (sh.image.ext or "png").lstrip(".")
    except Exception:
        return
    dest, ref = util.save_asset(blob, ctx, ext)
    if dest:
        items.append(f"![图{no}]({ref})\n\n<!-- ocr:img:{dest} -->")
    else:
        items.append(defer(deferred, "image-ocr", 1, "图片未存盘"))
