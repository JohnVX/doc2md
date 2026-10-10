"""Word(docx)解析: python-docx, 保序输出段落/表格 + 内嵌图提取+OCR占位 + 文本框 + 公式/图表标记.

通用, 不针对具体文档:
  - body 元素按序遍历 (w:p / w:tbl), 段落按样式映射为 md 标题/有序列表(编号+层级)/正文(含粗斜体)
  - 表格 -> md 表格; 表格单元格内的图片也提取
  - 段落内嵌图(含纯图段落) -> 提取 blob -> 存 assets -> OCR 占位(收尾执行)
  - 浮动文本框 (w:txbxContent) -> 扫描 XML 提取文字附在末尾
  - 公式(OMML)/图表/SmartArt -> 检测到即标记 stage2(文本框不计, 其文字已提取)
"""
import re
from pathlib import Path

from .. import util
from ._common import defer, is_placeholder, md_table


def parse(path, ctx=None):
    """docx 解析入口(段落/表格/有序列表/图片/公式图表)."""
    import docx
    from docx.oxml.ns import qn
    d = docx.Document(path)
    stem = Path(path).stem

    core_title = None
    try:
        core_title = d.core_properties.title
    except Exception:
        pass
    title = util.pick_title([core_title], stem)

    para_map = {p._element: p for p in d.paragraphs}
    tbl_map = {t._element: t for t in d.tables}
    body_el = d.element.body

    out = []
    deferred = []
    for child in body_el.iterchildren():
        if child in para_map:
            line = _render_paragraph(para_map[child], d, ctx, deferred)
            if line is not None:
                out.append(line)
        elif child in tbl_map:
            out.append(_render_table(tbl_map[child]))
            # 表格单元格内的图片(表格 md 放不下图, 单独附后提取+OCR, 不静默丢失)
            for i, (blob, ext) in enumerate(_element_images(child, d), 1):
                dest, ref = util.save_asset(blob, ctx, ext)
                out.append(f"![表格图{i}]({ref})\n\n<!-- ocr:img:{dest} -->")

    # 浮动文本框防御 (body 之外的 drawing 里的文本)
    txbx = _extract_textboxes(d)
    if txbx:
        out.append("### 浮动文本框内容")
        out.extend(txbx)
    # 诚实标记: 检测到但未提取的内容(公式/图表/SmartArt), 留给 stage2
    _detect_unextracted(d, deferred, out)

    body = "\n\n".join(x for x in out if x) or f"<!-- 空文档: {Path(path).name} -->"
    return {"title": title, "body": body, "tags": [], "meta": {}, "deferred": deferred}


def _render_paragraph(p, d, ctx, deferred):
    """渲染 docx 段落为 md."""
    style = (p.style.name if p.style else "") or ""
    sname = style.lower()
    text = _runs_md(p)

    # 标题样式
    m = re.match(r"heading\s*([1-9])", sname)
    if m:
        lvl = int(m.group(1))
        if text.strip():
            return "#" * lvl + " " + text.strip()
        return None
    if sname == "title" and text.strip():
        return "# " + text.strip()
    # 列表: 恢复有序编号与层级缩进
    # _list_info 从 numPr 推断(有 numId 时可靠); 无 numPr 时按样式名兜底(number/bullet)
    li = _list_info(p, d)
    if "list" in sname or li is not None:
        if text.strip():
            ordered = (li[0] if li else False) or ("number" in sname)
            level = li[1] if li else 0
            marker = "1." if ordered else "-"
            return "  " * level + marker + " " + text.strip()
        return None
    # 内嵌图(在空段判断前, 否则无文字的纯图段落会被跳过致图丢失)
    imgs = _para_images(p, d)
    if imgs:
        chunks = [text] if text.strip() else []
        for i, (blob, ext) in enumerate(imgs, 1):
            dest, ref = util.save_asset(blob, ctx, ext)
            chunks.append(f"![图{i}]({ref})\n\n<!-- ocr:img:{dest} -->")
        return "\n\n".join(chunks)
    # 空段
    if not text.strip():
        return ""
    return text


def _runs_md(p):
    """提取 run 级 markdown 文本."""
    # 合并相邻同样式(粗/斜)的 run, 避免产生 **** 杂讯
    segs = []  # [[bold, italic, text]]
    for r in p.runs:
        t = r.text or ""
        if not t:
            continue
        b, i = bool(r.bold), bool(r.italic)
        if segs and segs[-1][0] == b and segs[-1][1] == i:
            segs[-1][2] += t
        else:
            segs.append([b, i, t])
    out = []
    for b, i, t in segs:
        if not t:
            continue
        if b and i:
            out.append(f"***{t}***")
        elif b:
            out.append(f"**{t}**")
        elif i:
            out.append(f"*{t}*")
        else:
            out.append(t)
    return "".join(out) if out else (p.text or "")


def _list_info(p, d):
    """返回 (ordered, level) 或 None. 从 numPr + numbering part 推断有序/层级."""
    from docx.oxml.ns import qn
    pPr = p._element.find(qn("w:pPr"))
    if pPr is None:
        return None
    numPr = pPr.find(qn("w:numPr"))
    if numPr is None:
        return None
    ilvl_el = numPr.find(qn("w:ilvl"))
    numId_el = numPr.find(qn("w:numId"))
    level = 0
    if ilvl_el is not None:
        v = ilvl_el.get(qn("w:val"))
        if v and v.isdigit():
            level = int(v)
    ordered = False
    numId = numId_el.get(qn("w:val")) if numId_el is not None else None
    if numId:
        try:
            npart = getattr(d.part, "numbering_part", None)
            if npart is not None:
                root = npart.element
                abstract_id = None
                for num in root.findall(qn("w:num")):
                    if num.get(qn("w:numId")) == numId:
                        a = num.find(qn("w:abstractNumId"))
                        if a is not None:
                            abstract_id = a.get(qn("w:val"))
                        break
                if abstract_id is not None:
                    for an in root.findall(qn("w:abstractNum")):
                        if an.get(qn("w:abstractNumId")) == abstract_id:
                            for lvl in an.findall(qn("w:lvl")):
                                if lvl.get(qn("w:ilvl")) == str(level):
                                    fmt = lvl.find(qn("w:numFmt"))
                                    if fmt is not None:
                                        fv = fmt.get(qn("w:val")) or ""
                                        ordered = fv not in ("bullet", "none")
                                    break
                            break
        except Exception:
            pass
    return (ordered, level)


def _has_numpr(p):
    """检测段落是否有编号属性."""
    from docx.oxml.ns import qn
    pPr = p._element.find(qn("w:pPr"))
    if pPr is None:
        return False
    return pPr.find(qn("w:numPr")) is not None


def _element_images(element, d):
    """扫描任意元素子树内的 a:blip, 返回 [(blob, ext), ...]. 用于段落/表格内嵌图."""
    from docx.oxml.ns import qn
    out = []
    for blip in element.findall(".//" + qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        try:
            try:
                part = d.part.related_parts[rid]
            except Exception:
                part = d.part.rels[rid].target_part
            blob = part.blob
            ext = (part.partname.ext or "png").lstrip(".")
            out.append((blob, ext))
        except Exception:
            continue
    return out


def _para_images(p, d):
    """返回段落内嵌图的 [(blob, ext), ...]."""
    return _element_images(p._element, d)


def _render_table(tbl):
    """渲染 docx 表格为 md 表格."""
    try:
        ncols = len(tbl.columns)
    except Exception:
        ncols = max((len(r.cells) for r in tbl.rows), default=0)
    if ncols < 1 or not tbl.rows:
        return ""
    rows = []
    for r in tbl.rows:
        cells = [c.text.replace("\n", " ").replace("|", "\\|").strip() for c in r.cells]
        rows.append(cells)
    return md_table(rows, ncols)


def _extract_textboxes(d):
    """提取文本框内容."""
    from docx.oxml.ns import qn
    out = []
    for tb in d.element.iter(qn("w:txbxContent")):
        s = "".join((t.text or "") for t in tb.iter(qn("w:t"))).strip()
        if s and not is_placeholder(s):
            out.append(s)
    return out


def _detect_unextracted(d, deferred, out):
    """检测到但 stage1 未提取的内容(公式/图表/SmartArt), 追加 defer 标记到 out."""
    from docx.oxml.ns import qn
    # 数学公式 (OMML)
    try:
        n_eq = len(d.element.findall(".//" + qn("m:oMath")))
    except Exception:
        n_eq = 0
    if n_eq:
        out.append(defer(deferred, "formula", n_eq, "含数学公式(OMML)未提取"))
    # 非位图 drawing (图表/SmartArt 等): drawing 内无 a:blip 即非图片.
    # 排除文本框(含 w:txbxContent)——其文字已由 _extract_textboxes 提取, 不算未提取图形.
    try:
        txbx_qn = qn("w:txbxContent")
        blip_qn = qn("a:blip")
        n_graphic = sum(
            1 for dr in d.element.iter(qn("w:drawing"))
            if not dr.findall(".//" + blip_qn)
            and not dr.findall(".//" + txbx_qn)
        )
    except Exception:
        n_graphic = 0
    if n_graphic:
        out.append(defer(deferred, "chart", n_graphic, "图表/SmartArt等图形未提取"))
