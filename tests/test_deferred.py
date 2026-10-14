"""deferred 结构化数据测试: stage2 标记可被编程统一处理.

验证:
  1. front-matter 含 deferred 字段 (结构化: type/count/note)
  2. catalog 每文档有 deferred 上下文; stage2_pending 在 handoff 不在 catalog
  3. 标记格式统一: <!-- stage2:type=count note --> 可用正则提取
  4. 无 defer 的文档不输出 deferred 字段(不噪声)
  5. 多个同类型 defer 项被聚合(count 求和)
  6. handoff.yaml 完整: producer/契约/输出指南/stage2_pending
"""
import os
import re
import yaml

from tests._runner import (assert_eq, assert_true, cleanup, find_mds,
                            make_tmp_input, make_tmp_output, parse_front_matter,
                            read_md)

from doc2md import pipeline
from doc2md.parsers._common import DEFER_TYPES, aggregate_deferred, defer

CASES = []

_STAGE2_RE = re.compile(r'<!-- stage2:([\w-]+)=(\d+)(?:\s+(.*?))? -->')


def test_unsupported_has_structured_deferred():
    """非空不支持格式(ole2) -> front-matter + catalog 都有结构化 deferred."""
    in_dir = make_tmp_input({
        "note.txt": "正常文本\n",
        "old.doc": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1\x00" + b"\x00" * 100,
    })
    out = make_tmp_output()
    pipeline.run(in_dir, out)
    mds = find_mds(out)

    ole2_md = None
    txt_md = None
    for m in mds:
        fm = parse_front_matter(open(m, encoding="utf-8").read())
        if fm.get("source_type") == "ole2":
            ole2_md = m
        elif fm.get("source_type") == "txt":
            txt_md = m
    assert_true(ole2_md is not None, "ole2 文档应有 md")
    assert_true(txt_md is not None, "txt 文档应有 md")

    body = open(ole2_md, encoding="utf-8").read()
    fm = parse_front_matter(body)
    assert_true("deferred" in fm, "ole2 front-matter 应有 deferred 字段")
    assert_eq(fm["deferred"][0]["type"], "unsupported")
    assert_eq(fm["deferred"][0]["count"], 1)

    txt_body = open(txt_md, encoding="utf-8").read()
    txt_fm = parse_front_matter(txt_body)
    assert_true("deferred" not in txt_fm, "txt 无 defer, 不应输出 deferred 字段")

    matches = _STAGE2_RE.findall(body)
    assert_eq(len(matches), 1, "body 应有 1 个标准化 stage2 标记")
    assert_eq(matches[0][0], "unsupported", "标记 type=unsupported")
    assert_eq(matches[0][1], "1", "标记 count=1")

    cleanup(in_dir, out)


CASES.append(("unsupported_has_structured_deferred", test_unsupported_has_structured_deferred))


def test_stage2_pending_in_handoff_not_catalog():
    """stage2_pending 在 handoff.yaml, 不在 catalog.yaml."""
    in_dir = make_tmp_input({
        "ok.txt": "正常\n",
        "bad.doc": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1\x00" * 20,
    })
    out = make_tmp_output()
    pipeline.run(in_dir, out)

    cat = yaml.safe_load(open(os.path.join(out, "index", "catalog.yaml"), encoding="utf-8"))
    hdata = yaml.safe_load(open(os.path.join(out, "index", "handoff.yaml"), encoding="utf-8"))

    # catalog 不含 stage2_pending
    assert_true("stage2_pending" not in cat, "catalog 不应有 stage2_pending")

    # catalog 每文档 deferred 仍在 (上下文)
    for d in cat["documents"]:
        if d["source_type"] == "txt":
            assert_true("deferred" not in d, "txt 文档不应有 deferred")
        else:
            assert_true("deferred" in d, "ole2 文档应有 deferred 上下文")

    # handoff 有 stage2_pending
    assert_true("stage2_pending" in hdata, "handoff 应有 stage2_pending")
    assert_eq(len(hdata["stage2_pending"]), 1, "应有 1 个待 stage2 文档")
    entry = hdata["stage2_pending"][0]
    assert_eq(entry["deferred"][0]["type"], "unsupported")
    assert_true("doc" in entry and "id" in entry, "工作项应有 id 和 doc 路径")

    cleanup(in_dir, out)


CASES.append(("stage2_pending_in_handoff_not_catalog", test_stage2_pending_in_handoff_not_catalog))


def test_marker_format_parseable():
    """所有 stage2 标记都能用统一正则提取 type/count/note."""
    deferred = []
    markers = [
        defer(deferred, "image-ocr", 1, "无引擎或纯图无文字"),
        defer(deferred, "formula", 3, "含数学公式(OMML)未提取"),
        defer(deferred, "chart", 2, "图表/SmartArt等未提取"),
        defer(deferred, "encrypted", 1, "PDF加密"),
        defer(deferred, "unsupported", 1, "格式ole2"),
        defer(deferred, "parse-failed", 1, "RuntimeError"),
        defer(deferred, "formula-cell", 5, "无缓存值"),
        defer(deferred, "inline-image", 3, "图内文字未提取"),
        defer(deferred, "scan-page", 1, "OCR无结果"),
    ]
    body = "\n".join(markers)
    matches = _STAGE2_RE.findall(body)
    assert_eq(len(matches), len(markers), "所有标记都应被正则提取")
    types = [m[0] for m in matches]
    for t in types:
        assert_true(t in DEFER_TYPES, f"类型 {t} 应在 DEFER_TYPES 枚举内")
    assert_eq(matches[1][1], "3", "formula count=3")
    assert_eq(matches[3][1], "1", "encrypted count=1")
    m2 = defer(deferred, "image-ocr", 1)
    matches2 = _STAGE2_RE.findall(m2)
    assert_eq(len(matches2), 1, "无 note 的标记也应可提取")
    assert_eq(matches2[0][2], "", "无 note 时 note 为空")


CASES.append(("marker_format_parseable", test_marker_format_parseable))


def test_aggregate_deferred():
    """同类型 defer 项聚合: count 求和, note 取首个."""
    items = [
        {"type": "image-ocr", "count": 2, "note": "图1"},
        {"type": "image-ocr", "count": 3, "note": "图2"},
        {"type": "formula", "count": 1, "note": "公式1"},
    ]
    result = aggregate_deferred(items)
    by_type = {r["type"]: r for r in result}
    assert_eq(by_type["image-ocr"]["count"], 5, "image-ocr count 应聚合为 5")
    assert_eq(by_type["formula"]["count"], 1)
    assert_eq(len(result), 2, "聚合后应只有 2 种类型")


CASES.append(("aggregate_deferred", test_aggregate_deferred))


def test_no_deferred_for_clean_docs():
    """纯文本无 defer -> catalog 不带 deferred; handoff stage2_pending 为空."""
    in_dir = make_tmp_input({"clean.txt": "just text\n"})
    out = make_tmp_output()
    pipeline.run(in_dir, out)
    cat = yaml.safe_load(open(os.path.join(out, "index", "catalog.yaml"), encoding="utf-8"))
    hdata = yaml.safe_load(open(os.path.join(out, "index", "handoff.yaml"), encoding="utf-8"))
    assert_true("stage2_pending" not in cat, "catalog 不应有 stage2_pending")
    assert_true("deferred" not in cat["documents"][0], "clean doc 不应有 deferred")
    assert_eq(hdata["stage2_pending"], [], "无 defer 时 handoff stage2_pending 应为空列表")
    cleanup(in_dir, out)


CASES.append(("no_deferred_for_clean_docs", test_no_deferred_for_clean_docs))


def test_handoff_yaml_complete():
    """handoff.yaml 完整: producer/契约/输出指南/layout/stage2_pending."""
    in_dir = make_tmp_input({"x.txt": "text\n"})
    out = make_tmp_output()
    pipeline.run(in_dir, out)
    hpath = os.path.join(out, "index", "handoff.yaml")
    assert_true(os.path.isfile(hpath), "handoff.yaml 应存在")
    data = yaml.safe_load(open(hpath, encoding="utf-8"))
    assert_eq(data["producer"]["name"], "doc2md")
    assert_eq(data["producer"]["stage"], 1)
    assert_true(len(data["stage1_done"]) >= 5, "stage1_done 应列出多项能力")
    assert_true(len(data["stage1_not_done"]) >= 2, "stage1_not_done 应列出未做项")
    contract = data["deferred_contract"]
    for t in ("image-ocr", "inline-image", "scan-page", "formula", "chart",
              "formula-cell", "encrypted", "unsupported", "parse-failed"):
        assert_true(t in contract, f"deferred_contract 应含类型 {t}")
        assert_true("where" in contract[t] and "source" in contract[t] and "action" in contract[t],
                     f"类型 {t} 应有 where/source/action")
    assert_true("procedure" in data, "handoff 应有 procedure 操作步骤")
    assert_true(len(data["procedure"]) >= 5, "procedure 应至少 5 步")
    guide = data["output_guide"]
    assert_true("md_body" in guide and "front_matter" in guide
                and "catalog" in guide and "handoff" in guide,
                "output_guide 应含 md_body/front_matter/catalog/handoff 四处更新目标")
    assert_true("layout" in data and "docs" in data["layout"])
    assert_true("stage2_pending" in data, "handoff 应有 stage2_pending 字段(即使空)")
    cleanup(in_dir, out)


CASES.append(("handoff_yaml_complete", test_handoff_yaml_complete))


def test_catalog_has_audience_header():
    """catalog.yaml 顶部有面向对象注释, 引导 stage2 agent 读 handoff."""
    in_dir = make_tmp_input({"x.txt": "text\n"})
    out = make_tmp_output()
    pipeline.run(in_dir, out)
    txt = open(os.path.join(out, "index", "catalog.yaml"), encoding="utf-8").read()
    assert_true("面向" in txt, "catalog.yaml 应有面向对象注释")
    assert_true("handoff" in txt, "catalog 应引导 stage2 agent 读 handoff.yaml")
    cleanup(in_dir, out)


CASES.append(("catalog_has_audience_header", test_catalog_has_audience_header))


def test_aggregate_preserves_locs():
    """同类型多项的 loc 汇总为 locs 列表(去重), 不丢位置; 无 loc 时不产字段."""
    d = []
    defer(d, "formula", 1, "OMML", loc="p3")
    defer(d, "formula", 1, "OMML", loc="p7")
    defer(d, "formula", 1, "OMML", loc="p7")  # 重复应去重
    agg = aggregate_deferred(d)
    assert_eq(len(agg), 1)
    assert_eq(agg[0]["count"], 3)
    assert_eq(agg[0]["locs"], ["p3", "p7"])
    # 无 loc 时不应出现 locs 字段
    d2 = []
    defer(d2, "image-ocr", 1, "")
    assert_true("locs" not in aggregate_deferred(d2)[0], "无 loc 不应产出 locs 字段")


CASES.append(("aggregate_preserves_locs", test_aggregate_preserves_locs))


def test_docx_formula_marked_in_place():
    """docx 公式标记应紧随其所在段落(原位), 而非文档末尾汇总."""
    from docx import Document
    from docx.oxml.ns import qn
    in_dir = make_tmp_input()
    d = Document()
    d.add_paragraph("PARA_A")
    p1 = d.add_paragraph()
    p1._element.append(p1._element.makeelement(qn("m:oMath"), {}))
    d.add_paragraph("PARA_B")
    p2 = d.add_paragraph()
    p2._element.append(p2._element.makeelement(qn("m:oMath"), {}))
    p2._element.append(p2._element.makeelement(qn("m:oMath"), {}))
    d.add_paragraph("PARA_C")
    d.save(os.path.join(in_dir, "fm.docx"))
    out = make_tmp_output()
    _, _, e = pipeline.run(in_dir, out)
    assert_eq(e, 0)
    assert_eq(len(find_mds(out)), 1)
    txt = read_md(out)
    # 段1 的标记应在 PARA_A 与 PARA_B 之间
    ia, ib = txt.index("PARA_A"), txt.index("PARA_B")
    m1 = txt.index("stage2:formula=1")
    assert_true(ia < m1 < ib, "段1 公式标记应位于 PARA_A 之后 PARA_B 之前")
    # 段2 的两个公式合并为 formula=2, 位于 PARA_B 与 PARA_C 之间
    ic = txt.index("PARA_C")
    m2 = txt.index("stage2:formula=2")
    assert_true(ib < m2 < ic, "段2 公式标记应在 PARA_B 之后 PARA_C 之前")
    cleanup(in_dir, out)


CASES.append(("docx_formula_marked_in_place", test_docx_formula_marked_in_place))


def test_contract_lists_all_types():
    """handoff 契约必须覆盖全部 DEFER_TYPES, 且含终止条件与 loc 说明."""
    in_dir = make_tmp_input({"x.txt": "text\n"})
    out = make_tmp_output()
    pipeline.run(in_dir, out)
    data = yaml.safe_load(open(os.path.join(out, "index", "handoff.yaml"), encoding="utf-8"))
    contract = set(data["deferred_contract"])
    assert_eq(contract, set(DEFER_TYPES), "契约类型应与 DEFER_TYPES 完全一致")
    assert_true("termination" in data, "handoff 应有终止条件说明")
    assert_true("blocked" in data["termination"], "终止条件应定义 blocked 退路")
    assert_true("loc_field" in data, "handoff 应说明 locs 字段")
    cleanup(in_dir, out)


CASES.append(("contract_lists_all_types", test_contract_lists_all_types))
