"""回归测试: 通用性/健壮性审查中修复的缺陷.

覆盖:
- is_placeholder 整串匹配 (不误删含 "title"/"click to add" 的真实正文)
- is_generic_title 不再因含 "xxx" 子串误杀真实标题
- 解析器在无 output_root 时不产出 (None) 引用, 降级为 defer
- pptx 标题占位符不重复输出; 多行文本框不当标题(不丢行)
- heading_slugs 对重复标题生成唯一锚点
"""
import os
import tempfile

from tests._runner import assert_eq, assert_true, make_tmp_input, make_tmp_output, cleanup

CASES = []


def test_placeholder_exact_match():
    """占位符整串匹配: 样板命中, 真实正文(含关键词)保留."""
    from doc2md.parsers._common import is_placeholder
    for t in ("单击此处添加标题", "占位符", "Click to add text", "Click to add title",
              "Placeholder text", "Title", "Title 1"):
        assert_true(is_placeholder(t), f"{t!r} 应判为占位符")
    for t in ("Set the title style to bold", "My Title Here",
              "The title of the paper is X", "Click here to add a note",
              "A user may click to add items to the cart", "幻灯片标题AA"):
        assert_true(not is_placeholder(t), f"{t!r} 是真实正文, 不应判为占位符")


def test_generic_title_not_substring_xxx():
    """含 xxx 子串的真实标题不应判为通用."""
    from doc2md.util import is_generic_title
    assert_true(is_generic_title("XXX"), "纯 XXX 应判通用")
    assert_true(is_generic_title("xxx"), "纯 xxx 应判通用")
    assert_true(is_generic_title(""), "空标题应判通用")
    for t in ("Report XXX-2024", "保险条款xxx版", "XXX项目计划", "MyDocument"):
        assert_true(not is_generic_title(t), f"{t!r} 是真实标题, 不应判通用")


def test_parser_without_output_root_defers_not_none_ref():
    """无 output_root 时: 不产出 ![](None), 降级为 image-ocr defer."""
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    from doc2md.parsers import pptx as pptxmod
    with tempfile.TemporaryDirectory() as td:
        imgp = os.path.join(td, "x.png")
        Image.new("RGB", (8, 8), (0, 0, 0)).save(imgp)
        path = os.path.join(td, "pic.pptx")
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(imgp, Inches(1), Inches(1))
        prs.save(path)
        r = pptxmod.parse(path, ctx={})
    assert_true("(None)" not in r["body"], "不应出现 ![](None) 引用")
    assert_true("ocr:img:None" not in r["body"], "不应出现 ocr:img:None 占位")
    types = {d["type"] for d in r.get("deferred", [])}
    assert_true("image-ocr" in types, "应降级为 image-ocr defer")


def test_pptx_title_placeholder_not_duplicated():
    """有 title 占位符时, 标题只出现在 heading, 不在正文重复."""
    from pptx import Presentation
    from doc2md.parsers import pptx as pptxmod
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "t.pptx")
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
        s.shapes.title.text = "唯一标题XYZ"
        prs.save(path)
        r = pptxmod.parse(path, ctx={})
    assert_eq(r["body"].count("唯一标题XYZ"), 1, "标题不应重复")


def test_pptx_multiline_textbox_kept():
    """无占位符时多行文本框不当标题, 全部行保留 (不丢内容)."""
    from pptx import Presentation
    from pptx.util import Inches
    from doc2md.parsers import pptx as pptxmod
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "m.pptx")
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
        tb.text_frame.text = "第一行标题\n第二行内容\n第三行内容"
        prs.save(path)
        r = pptxmod.parse(path, ctx={})
    for ln in ("第一行标题", "第二行内容", "第三行内容"):
        assert_true(ln in r["body"], f"多行文本框内容 {ln!r} 应保留")


def test_pptx_real_content_with_title_word_kept():
    """含 'title' 一词的真实正文不应被占位符过滤丢弃."""
    from pptx import Presentation
    from pptx.util import Inches
    from doc2md.parsers import pptx as pptxmod
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "c.pptx")
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        a = s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        a.text_frame.text = "幻灯片标题AA"
        b = s.shapes.add_textbox(Inches(1), Inches(3), Inches(5), Inches(1))
        b.text_frame.text = "Set the title style to bold"
        prs.save(path)
        r = pptxmod.parse(path, ctx={})
    assert_true("Set the title style to bold" in r["body"], "含 title 的真实正文应保留")


def test_heading_slugs_unique():
    """重复标题生成唯一锚点; 连字符保留."""
    from doc2md.util import heading_slugs, heading_slug
    assert_eq(heading_slugs(["安装说明", "安装说明", "安装说明"]),
              ["安装说明", "安装说明-1", "安装说明-2"], "重复标题锚点应加后缀")
    assert_eq(heading_slug("a-b"), "a-b", "连字符应保留")
    assert_eq(heading_slugs(["A", "B"]), ["a", "b"], "无重复时用原名")


CASES.append(("placeholder_exact_match", test_placeholder_exact_match))
CASES.append(("generic_title_not_substring_xxx", test_generic_title_not_substring_xxx))
CASES.append(("parser_without_output_root_defers_not_none_ref",
              test_parser_without_output_root_defers_not_none_ref))
CASES.append(("pptx_title_placeholder_not_duplicated",
              test_pptx_title_placeholder_not_duplicated))
CASES.append(("pptx_multiline_textbox_kept", test_pptx_multiline_textbox_kept))
CASES.append(("pptx_real_content_with_title_word_kept",
              test_pptx_real_content_with_title_word_kept))
CASES.append(("heading_slugs_unique", test_heading_slugs_unique))
