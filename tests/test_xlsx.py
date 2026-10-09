"""xlsx 测试: 无缓存值的公式单元格应显示公式文本(而非空)."""
import os

from tests._runner import (assert_true, cleanup, make_tmp_input, make_tmp_output,
                            read_md)

from doc2md import pipeline
from doc2md.parsers import xlsx as xlsx_parser

CASES = []


def test_formula_without_cached_shown():
    """脚本生成的 xlsx(公式无缓存) -> md 应含公式文本, 不为空."""
    import openpyxl
    in_dir = make_tmp_input()
    p = os.path.join(in_dir, "f.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "项", "值"
    ws["A2"], ws["B2"] = "x", '=A2&"-f"'   # 公式, 无缓存
    ws["A3"], ws["B3"] = 1, "=A3*2"
    wb.save(p)

    body = xlsx_parser.parse(p)["body"]
    assert_true("=A3*2" in body, "应显示公式文本 =A3*2")
    assert_true("=A2" in body, "应显示公式文本 =A2&...")

    out = make_tmp_output()
    pipeline.run(in_dir, out)
    txt = read_md(out)
    assert_true("=A3*2" in txt)
    assert_true("stage2" in txt, "应有公式回退标记")
    cleanup(in_dir, out)


CASES.append(("formula_without_cached_shown", test_formula_without_cached_shown))


def test_cached_values_normal():
    """纯值(无公式) -> 不应出现公式回退标记."""
    import openpyxl
    in_dir = make_tmp_input()
    p = os.path.join(in_dir, "v.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "k", "v"
    ws["A2"], ws["B2"] = "a", "b"
    wb.save(p)
    body = xlsx_parser.parse(p)["body"]
    assert_true("=A" not in body, "无公式时不应出现公式回退标记")
    assert_true("stage2" not in body)
    cleanup(in_dir)


CASES.append(("cached_values_normal", test_cached_values_normal))


def test_merged_cell_expanded():
    """合并单元格的值应填满跨度(非仅左上角)."""
    import openpyxl
    in_dir = make_tmp_input()
    p = os.path.join(in_dir, "merged.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"] = "维度", "子维度", "值"
    ws["A2"] = "大类X"
    ws.merge_cells("A2:A4")  # 垂直合并 A2:A4
    ws["B2"], ws["C2"] = "子1", "v1"
    ws["B3"], ws["C3"] = "子2", "v2"
    ws["B4"], ws["C4"] = "子3", "v3"
    wb.save(p)
    out = make_tmp_output()
    pipeline.run(in_dir, out)
    body = read_md(out)
    # A3, A4 应回填"大类X"(垂直合并展开), 而非空
    lines = body.splitlines()
    a3_a4 = [ln for ln in lines if "大类X" in ln]
    assert_true(len(a3_a4) >= 3, f"合并区 A2:A4 三行都应有'大类X', 实得含'大类X'行数={len(a3_a4)}")
    cleanup(in_dir, out)


CASES.append(("merged_cell_expanded", test_merged_cell_expanded))
