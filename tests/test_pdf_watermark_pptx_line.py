"""回归测试: PDF 跨页样板(页眉/页脚/水印)跳过 + 表格单元格水印污染清洗 + PPTX 箭头不计入 defer.

对应实际产物检查中的发现:
- PDF 每页重复的页眉水印污染正文, 且与表格 bbox 重叠时把水印字符混入单元格
  (如 g_aaaBbb -> "g aaaBbb _", 下划线错位)
- PPTX 里的连接符(LINE/箭头)被计入 chart defer, 导致 defer 计数虚高
"""
import os
import tempfile

from tests._runner import assert_eq, assert_true

from doc2md.parsers import pdf as pdfmod

CASES = []


def _make_watermark_pdf(path):
    """生成 5 页 PDF: 每页同位置同文字的页眉水印 + 各页不同正文."""
    import pymupdf as fitz
    doc = fitz.open()
    for i in range(5):
        p = doc.new_page()
        p.insert_text((72, 72), f"page {i + 1} unique body text")
        p.insert_text((200, 300), "WATERMARK confidential")
    doc.save(path)
    doc.close()


def test_pdf_boilerplate_detects_watermark():
    """跨页重复的页眉水印应被检出为样板."""
    import pymupdf as fitz
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "wm.pdf")
        _make_watermark_pdf(path)
        d = fitz.open(path)
        bp = pdfmod._find_boilerplate(d)
        d.close()
    assert_true(any("WATERMARK" in t for t, *_ in bp), "水印应被检出")
    assert_true(not any("unique body" in t for t, *_ in bp), "正文不应被误判为样板")


def test_pdf_watermark_removed_from_body():
    """端到端: 水印不出现在产物正文, 正文保留."""
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "wm.pdf")
        _make_watermark_pdf(path)
        res = pdfmod.parse(path)
    body = res["body"]
    assert_true("WATERMARK" not in body, "水印应被剔除")
    assert_true("unique body text" in body, "正文应保留")


def test_pdf_no_false_boilerplate_on_plain_pdf():
    """普通 PDF (无重复样板) 不应误判."""
    import pymupdf as fitz
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "plain.pdf")
        doc = fitz.open()
        for i in range(5):
            p = doc.new_page()
            p.insert_text((72, 72), f"unique line number {i} abc")
        doc.save(path)
        doc.close()
        d = fitz.open(path)
        bp = pdfmod._find_boilerplate(d)
        d.close()
    assert_eq(len(bp), 0, "无重复样板时应检出 0")


def test_pptx_connector_not_counted_as_chart_defer():
    """PPTX 里只有连接符(箭头)时, 不应产生 chart defer."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Inches
    from doc2md.parsers import pptx as pptxmod
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "line.pptx")
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                               Inches(1), Inches(1), Inches(3), Inches(3))
        prs.save(path)
        res = pptxmod.parse(path)
    types = {d["type"] for d in res.get("deferred", [])}
    assert_true("chart" not in types, "纯连接符不应产生 chart defer")


def test_pdf_boilerplate_only_page_not_scanned():
    """仅含页眉/水印的页 (剔除样板后变空) 不应被当作扫描页."""
    import pymupdf as fitz
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "hdr.pdf")
        doc = fitz.open()
        for i in range(5):
            p = doc.new_page()
            p.insert_text((72, 72), "HEADER repeated every page")
            if i == 2:
                p.insert_text((72, 200), "only real body here")
        doc.save(path)
        doc.close()
        res = pdfmod.parse(path)
    body = res["body"]
    assert_true("ocr:scan" not in body, "样板-only 页不应渲染为扫描页")
    assert_true("HEADER" not in body, "页眉样板应被剔除")
    assert_true("only real body here" in body, "正文应保留")


CASES.append(("pdf_boilerplate_detects_watermark", test_pdf_boilerplate_detects_watermark))
CASES.append(("pdf_watermark_removed_from_body", test_pdf_watermark_removed_from_body))
CASES.append(("pdf_no_false_boilerplate_on_plain_pdf",
              test_pdf_no_false_boilerplate_on_plain_pdf))
def test_cell_text_keeps_large_font():
    """大字号单元格文字不应被水印过滤误删 (相对高度判据)."""
    import pymupdf as fitz
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "big.pdf")
        doc = fitz.open()
        p = doc.new_page()
        p.insert_text((60, 100), "BigHeaderCell", fontsize=36)
        doc.save(path)
        doc.close()
        d = fitz.open(path)
        txt = pdfmod._cell_text(d[0], fitz.Rect(50, 60, 400, 120))
        d.close()
    assert_eq(txt, "BigHeaderCell", "大字号文字应保留")


CASES.append(("pdf_boilerplate_only_page_not_scanned",
              test_pdf_boilerplate_only_page_not_scanned))
CASES.append(("cell_text_keeps_large_font", test_cell_text_keeps_large_font))

CASES.append(("pptx_connector_not_counted_as_chart_defer",
              test_pptx_connector_not_counted_as_chart_defer))
