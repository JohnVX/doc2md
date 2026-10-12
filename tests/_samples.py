"""生成各格式样本文件 (hermetic: 测试输入全靠代码生成, 无外部 fixture 依赖).

用工程已有的库(python-docx/openpyxl/python-pptx/PyMuPDF/Pillow)现场造文件.
"""
import os


def gen_docx(path):
    from docx import Document
    d = Document()
    d.add_heading("样本标题", 0)
    d.add_paragraph("样本正文段落。")
    d.save(path)


def gen_xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "列1", "列2"
    ws.append(["a", "b"])
    ws.append(["c", "d"])
    wb.save(path)


def gen_pptx(path):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "样本幻灯片标题"
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = "样本正文"
    prs.save(path)


def gen_pdf(path):
    import pymupdf as fitz
    doc = fitz.open()
    p = doc.new_page()
    p.insert_text((50, 72), "Hello 样本 PDF")
    doc.save(path)
    doc.close()


def gen_png(path):
    from PIL import Image
    Image.new("RGB", (16, 16), (255, 0, 0)).save(path)


def gen_jpeg(path):
    from PIL import Image
    Image.new("RGB", (16, 16), (0, 255, 0)).save(path, "JPEG")


def gen_md(path):
    with open(path, "w", encoding="utf-8") as f:
        f.write("# 样本标题\n\n- 列表项一\n- 列表项二\n\n样本正文段落。\n")


def gen_code(path):
    with open(path, "w", encoding="utf-8") as f:
        f.write("cmake_minimum_required(VERSION 3.16)\nproject(x)\n")


def gen_txt(path):
    with open(path, "w", encoding="utf-8") as f:
        f.write("一段纯文本样本。\n")


def gen_pptx_rich(path):
    """PPTX with title, body, table, notes — covers real slide content."""
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    # Slide 1: Title + body
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "幻灯片标题"
    slide.placeholders[1].text = "正文内容段落"
    # Slide 2: Table
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "表格幻灯片"
    tbl = slide2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text = "表头A"
    tbl.cell(0, 1).text = "表头B"
    tbl.cell(1, 0).text = "数据1"
    tbl.cell(1, 1).text = "数据2"
    # Notes on slide 1
    slide.notes_slide.notes_text_frame.text = "这是备注内容"
    prs.save(path)


def gen_pdf_encrypted(path):
    """Encrypted PDF (user password) — should trigger `encrypted` defer."""
    import pymupdf as fitz
    doc = fitz.open()
    doc.new_page().insert_text((50, 72), "Secret content")
    doc.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="pass123")
    doc.close()


def gen_pptx_with_image(path, img_path=None):
    """PPTX with a slide image — for OCR/defer testing."""
    from pptx import Presentation
    from pptx.util import Inches
    import os
    if img_path is None:
        img_path = os.path.join(os.path.dirname(path), "_tmp_img.png")
        gen_text_png(img_path)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "带图幻灯片"
    slide.shapes.add_picture(img_path, Inches(1), Inches(2), Inches(4), Inches(2))
    prs.save(path)


def gen_xlsx_multisheet(path):
    """Multi-sheet xlsx — verify all sheets extracted."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "汇总"
    ws1["A1"] = "汇总数据"
    ws2 = wb.create_sheet("明细")
    ws2["A1"] = "明细数据"
    wb.save(path)


def gen_gbk_txt(path):
    """GBK-encoded Chinese text — verify encoding sniffing + correct decode."""
    with open(path, "wb") as f:
        f.write("中文GBK编码的纯文本。".encode("gbk"))


def _load_font(size):
    """跨平台尝试加载一个 TTF, 找不到用 PIL 默认."""
    from PIL import ImageFont
    for p in ("/usr/share/fonts/dejavu/DejaVuSans.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
              "/usr/share/fonts/dejavu-sans-fonts/DejaVuSans.ttf",
              "C:/Windows/Fonts/arial.ttf",
              "/System/Library/Fonts/Supplemental/Arial.ttf"):
        try:
            return ImageFont.truetype(p, size)
        except Exception:
            continue
    return ImageFont.load_default()


def gen_text_png(path, text="OCRTEST123"):
    """生成含清晰拉丁文字的图片(供 OCR 正向测试). DejaVu 含拉丁字形可被 OCR."""
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (320, 80), (255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((10, 20), text, fill=(0, 0, 0), font=_load_font(44))
    img.save(path)


_GEN = {
    "docx": gen_docx, "xlsx": gen_xlsx, "pptx": gen_pptx, "pdf": gen_pdf,
    "png": gen_png, "jpeg": gen_jpeg, "md": gen_md, "code": gen_code, "txt": gen_txt,
}

# kind -> detector 期望识别出的 format
_EXPECTED_FORMAT = {
    "docx": "docx", "xlsx": "xlsx", "pptx": "pptx", "pdf": "pdf",
    "png": "png", "jpeg": "jpeg", "md": "md", "code": "code", "txt": "txt",
}

_PLAN = [
    ("sample.docx", "docx"), ("data.xlsx", "xlsx"), ("slides.pptx", "pptx"),
    ("doc.pdf", "pdf"), ("pic.png", "png"), ("photo.jpg", "jpeg"),
    ("readme.md", "md"), ("CMakeLists.txt", "code"), ("note.txt", "txt"),
]


def generate_corpus(dest_dir):
    """在 dest_dir 生成全套样本, 返回 {filename: expected_format}."""
    expected = {}
    for name, kind in _PLAN:
        _GEN[kind](os.path.join(dest_dir, name))
        expected[name] = _EXPECTED_FORMAT[kind]
    return expected


def generate(dest_dir, kind, name):
    """生成单个指定 kind 的样本."""
    _GEN[kind](os.path.join(dest_dir, name))
    return _EXPECTED_FORMAT[kind]
