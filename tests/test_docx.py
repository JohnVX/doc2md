"""docx 解析测试: 表格单元格内的图片应被提取+OCR(回归 9 图全丢的 bug)."""
import os
import tempfile

from tests import _samples
from tests._runner import (HAS_RAPIDOCR, assert_contains_any, assert_eq,
                            assert_true, cleanup, find_md_by_source,
                            find_mds, make_tmp_input, make_tmp_output, read_md)

from doc2md import pipeline

CASES = []


def test_table_cell_image_extracted():
    # png 放输入目录外, 避免它被当独立图处理(只让 docx 作为输入)
    png_dir = tempfile.mkdtemp()
    png = os.path.join(png_dir, "cell.png")
    _samples.gen_text_png(png, "CELLOCR99")
    try:
        from docx import Document
        from docx.shared import Inches
        d = Document()
        d.add_paragraph("表格前的正文")
        table = d.add_table(rows=1, cols=1)
        table.cell(0, 0).paragraphs[0].add_run().add_picture(png, width=Inches(1))
        d.add_paragraph("表格后的正文")
        in_dir = make_tmp_input()
        d.save(os.path.join(in_dir, "with_table_img.docx"))

        out = make_tmp_output()
        _, _, e = pipeline.run(in_dir, out)
        assert_eq(e, 0)
        assert_eq(len(find_mds(out)), 1, "只应有一个 md(仅 docx 作为输入)")
        txt = find_md_by_source(out, "original-doc/with_table_img.docx")
        assert_true("表格图" in txt, "md 应含表格图标记")
        if HAS_RAPIDOCR:
            assert_contains_any(txt, ["cell", "ocr", "99"], "OCR 应识别表格图文字")
        cleanup(in_dir, out)
    finally:
        cleanup(png_dir)


CASES.append(("table_cell_image_extracted", test_table_cell_image_extracted))


def test_ordered_list_rendered():
    """有序列表(List Number 样式)应渲染为 1. 而非 -. 嵌套应有缩进."""
    from docx import Document
    in_dir = make_tmp_input()
    p = os.path.join(in_dir, "list.docx")
    d = Document()
    d.add_paragraph("步骤一", style="List Number")
    d.add_paragraph("步骤二", style="List Number")
    d.add_paragraph("普通段")
    d.save(p)
    out = make_tmp_output()
    _, _, e = pipeline.run(in_dir, out)
    assert_eq(e, 0)
    txt = read_md(out)
    assert_true("1. 步骤一" in txt, f"有序列表应渲染为 '1. 步骤一', 抽样: {txt[:200]!r}")
    assert_true("1. 步骤二" in txt, "第二个有序项也应为 1. (markdown 自动编号)")
    cleanup(in_dir, out)


CASES.append(("ordered_list_rendered", test_ordered_list_rendered))


def test_image_only_paragraph_extracted():
    """纯图段落(无文字)的图也应被提取(回归图被'空段'判断提前 return 跳过的 bug)."""
    png_dir = tempfile.mkdtemp()
    png = os.path.join(png_dir, "p.png")
    _samples.gen_text_png(png, "ONLYIMG77")
    try:
        from docx import Document
        from docx.shared import Inches
        d = Document()
        d.add_paragraph().add_run().add_picture(png, width=Inches(1))  # 纯图段落, 无文字
        d.add_paragraph("图后的正文")
        in_dir = make_tmp_input()
        d.save(os.path.join(in_dir, "img_only.docx"))
        out = make_tmp_output()
        _, _, e = pipeline.run(in_dir, out)
        assert_eq(e, 0)
        assert_eq(len(find_mds(out)), 1)
        txt = find_md_by_source(out, "original-doc/img_only.docx")
        assert_true("图1" in txt, f"纯图段落的图应提取, 抽样: {txt[:200]!r}")
        if HAS_RAPIDOCR:
            assert_contains_any(txt, ["only", "img", "77"], "纯图段落应 OCR")
        cleanup(in_dir, out)
    finally:
        cleanup(png_dir)


CASES.append(("image_only_paragraph_extracted", test_image_only_paragraph_extracted))


def test_footnote_endnote_comment_extracted():
    """脚注/尾注/批注(独立 part, python-docx 无 API)应被提取; 分隔符条目跳过."""
    import zipfile
    from docx import Document
    in_dir = make_tmp_input()
    p = os.path.join(in_dir, "notes.docx")
    d = Document()
    d.add_paragraph("正文段落")
    d.save(p)

    # 直接改写 zip: 追加 footnotes/endnotes/comments part + 正文引用
    data = {}
    with zipfile.ZipFile(p) as z:
        for n in z.namelist():
            data[n] = z.read(n)
    doc = data["word/document.xml"].decode("utf-8").replace(
        "</w:body>",
        '<w:p><w:r><w:footnoteReference w:id="1"/></w:r></w:p></w:body>')
    data["word/document.xml"] = doc.encode("utf-8")
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    data["word/footnotes.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:footnotes xmlns:w="{W}">'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1"><w:p><w:r><w:t>FN_MARKER_11</w:t></w:r></w:p></w:footnote>'
        '</w:footnotes>').encode("utf-8")
    data["word/endnotes.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:endnotes xmlns:w="{W}">'
        '<w:endnote w:id="1"><w:p><w:r><w:t>EN_MARKER_22</w:t></w:r></w:p></w:endnote>'
        '</w:endnotes>').encode("utf-8")
    data["word/comments.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:comments xmlns:w="{W}">'
        '<w:comment w:id="0" w:author="Reviewer"><w:p><w:r><w:t>CM_MARKER_33</w:t></w:r></w:p></w:comment>'
        '</w:comments>').encode("utf-8")
    ct = data["[Content_Types].xml"].decode("utf-8").replace(
        "</Types>",
        '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
        '<Override PartName="/word/endnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml"/>'
        '<Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>'
        "</Types>")
    data["[Content_Types].xml"] = ct.encode("utf-8")
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in data.items():
            z.writestr(n, b)

    out = make_tmp_output()
    _, _, e = pipeline.run(in_dir, out)
    assert_eq(e, 0, "run 应成功")
    txt = find_md_by_source(out, "original-doc/notes.docx")
    assert_true("FN_MARKER_11" in txt, "脚注应提取")
    assert_true("EN_MARKER_22" in txt, "尾注应提取")
    assert_true("CM_MARKER_33" in txt, "批注应提取")
    assert_true("Reviewer" in txt, "批注作者应标注")
    cleanup(in_dir, out)


CASES.append(("footnote_endnote_comment_extracted", test_footnote_endnote_comment_extracted))


def test_docx_without_notes_has_no_empty_section():
    """无脚注/尾注/批注时不产生空小节标题."""
    from docx import Document
    in_dir = make_tmp_input()
    d = Document()
    d.add_paragraph("只有正文")
    d.save(os.path.join(in_dir, "plain.docx"))
    out = make_tmp_output()
    _, _, e = pipeline.run(in_dir, out)
    assert_eq(e, 0)
    txt = find_md_by_source(out, "original-doc/plain.docx")
    for h in ("### 脚注", "### 尾注", "### 批注"):
        assert_true(h not in txt, f"无内容不应出现 {h}")
    cleanup(in_dir, out)


CASES.append(("docx_without_notes_no_empty_section", test_docx_without_notes_has_no_empty_section))
