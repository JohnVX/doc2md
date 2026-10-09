"""增量测试: 未变跳过 / 内容变重处理 / 删除清理 / 格式变 defer / 重分类剪枝."""
import os

import yaml

from tests import _samples
from tests._runner import (assert_eq, assert_true, cleanup, find_mds,
                            make_tmp_input, make_tmp_output)

from doc2md import pipeline

CASES = []

CFG = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "config.example.yaml")


def test_unchanged_skipped():
    in_dir = make_tmp_input()
    _samples.generate_corpus(in_dir)
    out = make_tmp_output()
    pipeline.run(in_dir, out)
    # 第二遍: 全部应跳过
    p, s, e = pipeline.run(in_dir, out)
    assert_eq(p, 0, "第二遍应无处理")
    assert_true(s >= 1, "应有跳过")
    assert_eq(e, 0)
    cleanup(in_dir, out)


CASES.append(("unchanged_skipped", test_unchanged_skipped))


def test_content_changed_reprocessed():
    in_dir = make_tmp_input({"note.txt": "原始内容\n"})
    out = make_tmp_output()
    pipeline.run(in_dir, out)
    # 改内容
    with open(os.path.join(in_dir, "note.txt"), "w", encoding="utf-8") as f:
        f.write("改动后的内容\n")
    p, s, e = pipeline.run(in_dir, out)
    assert_eq(p, 1, "内容变了应重处理1个")
    assert_eq(e, 0)
    # 新内容应在 md 里
    mds = find_mds(out)
    assert_true(any("改动后的内容" in open(f, encoding="utf-8").read() for f in mds))
    cleanup(in_dir, out)


CASES.append(("content_changed_reprocessed", test_content_changed_reprocessed))


def test_deleted_pruned():
    in_dir = make_tmp_input({"a.txt": "A\n", "b.txt": "B\n"})
    out = make_tmp_output()
    pipeline.run(in_dir, out)
    assert_eq(len(find_mds(out)), 2)
    os.remove(os.path.join(in_dir, "b.txt"))
    pipeline.run(in_dir, out)
    assert_eq(len(find_mds(out)), 1, "删除的文件其 md 应被清理")
    cleanup(in_dir, out)


CASES.append(("deleted_pruned", test_deleted_pruned))


def test_format_becomes_unsupported_deferred():
    """已处理文件内容换成非空不支持格式(含 null 二进制) -> 旧 md 被替换为 defer 标记."""
    in_dir = make_tmp_input({"note.txt": "正常文本\n"})
    out = make_tmp_output()
    pipeline.run(in_dir, out)
    assert_eq(len(find_mds(out)), 1)
    # 内容换成含 null 的二进制(非空, detector 判 binary)
    with open(os.path.join(in_dir, "note.txt"), "wb") as f:
        f.write(b"\x00" + os.urandom(511))  # 确定性含 null -> 识别为 binary
    pipeline.run(in_dir, out)
    mds = find_mds(out)
    assert_eq(len(mds), 1)
    assert_true("stage2" in open(mds[0], encoding="utf-8").read())
    man = yaml.safe_load(open(os.path.join(out, "index", ".manifest.yaml"), encoding="utf-8"))
    assert_eq(len(man["entries"]), 1, "manifest 仍保留该条目(defer)")
    cleanup(in_dir, out)


CASES.append(("format_becomes_unsupported_deferred", test_format_becomes_unsupported_deferred))


def test_reclassification_prunes_old_md():
    """内容变化导致分类变 -> 旧类目下旧 md 被删, 新类目下写新 md, 无孤儿."""
    in_dir = make_tmp_input({"note.txt": "一些中性内容文字\n"})  # 初内容无关键词 -> unclassified
    out = make_tmp_output()
    pipeline.run(in_dir, out, config_path=CFG)
    # 改内容使其命中 attack-surface 关键词(漏洞/利用/攻击)
    with open(os.path.join(in_dir, "note.txt"), "w", encoding="utf-8") as f:
        f.write("漏洞 利用 攻击面分析\n")
    pipeline.run(in_dir, out, config_path=CFG)
    mds = find_mds(out)
    # 同一 doc-id 的 md 只应在新类目(attack-surface)存在一份, 旧类目无孤儿
    assert_eq(len(mds), 1, "应只剩新类目下一份 md")
    assert_true("attack-surface" in mds[0], f"应在 attack-surface 目录, 实际: {mds[0]}")
    cat = yaml.safe_load(open(os.path.join(out, "index", "catalog.yaml"), encoding="utf-8"))
    assert_true("attack-surface" in cat["categories"])
    cleanup(in_dir, out)


CASES.append(("reclassification_prunes_old_md", test_reclassification_prunes_old_md))


def test_format_change_clears_old_assets():
    """支持格式(有 assets)变不支持(defer) -> 旧 assets 目录被清, 无孤儿."""
    import os as _os
    import tempfile
    from tests import _samples
    # 造一个含表格图的 docx(会提取 assets), 处理
    png_dir = tempfile.mkdtemp()
    png = _os.path.join(png_dir, "c.png")
    _samples.gen_text_png(png, "IMGASSET22")
    try:
        from docx import Document
        from docx.shared import Inches
        d = Document()
        d.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run().add_picture(png, width=Inches(1))
        in_dir = make_tmp_input()
        d.save(_os.path.join(in_dir, "with_img.docx"))
        out = make_tmp_output()
        pipeline.run(in_dir, out)
        # 应有 assets
        import glob as _g
        assert_true(len(_g.glob(_os.path.join(out, "original-doc", "*_assets", "*"))) >= 1,
                    "首次处理应有 assets")
        # 把内容换成含 null 的二进制 -> defer
        with open(_os.path.join(in_dir, "with_img.docx"), "wb") as f:
            f.write(b"\x00" + _os.urandom(511))
        pipeline.run(in_dir, out)
        # 旧 assets 应被清(格式变 defer, old_assets rmtree)
        assets_left = _g.glob(_os.path.join(out, "original-doc", "*_assets", "*"))
        assert_eq(len(assets_left), 0, f"旧 assets 应清, 残留: {assets_left}")
        cleanup(in_dir, out)
    finally:
        cleanup(png_dir)


CASES.append(("format_change_clears_old_assets", test_format_change_clears_old_assets))
