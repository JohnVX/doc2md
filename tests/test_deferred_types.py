"""deferred 类型覆盖: docx 公式/图表标记, 多 type 同文档, 标记位置正确性."""
import os

from tests._runner import (assert_eq, assert_true, cleanup, find_mds,
                            make_tmp_input, make_tmp_output, parse_front_matter,
                            require_ocr)

from doc2md import pipeline

CASES = []


def _make_docx_with_formula(path):
    """创建含 OMML 公式的 docx."""
    import docx
    from docx.oxml import OxmlElement
    d = docx.Document()
    d.add_paragraph("正文段落")
    para = d.add_paragraph()
    omath = OxmlElement('m:oMath')
    r = OxmlElement('m:r')
    t = OxmlElement('m:t')
    t.text = "E=mc2"
    r.append(t)
    omath.append(r)
    para._element.append(omath)
    d.save(path)


def test_docx_formula_deferred():
    """docx 含 OMML 公式 -> deferred type=formula, body 有标记."""
    in_dir = make_tmp_input({})
    p = os.path.join(in_dir, "formula.docx")
    _make_docx_with_formula(p)
    out = make_tmp_output()
    _, _, e = pipeline.run(in_dir, out)
    assert_eq(e, 0)
    body = open(find_mds(out)[0], encoding="utf-8").read()
    fm = parse_front_matter(body)
    assert_true("deferred" in fm, "应有 deferred")
    types = [d["type"] for d in fm["deferred"]]
    assert_true("formula" in types, f"应有 formula 类型, got {types}")
    assert_true("stage2:formula=" in body, "body 应有 formula 标记")
    cleanup(in_dir, out)


CASES.append(("docx_formula_deferred", test_docx_formula_deferred))


def _make_docx_with_chart(path):
    """创建含图表 drawing (无 a:blip) 的 docx."""
    import docx
    from docx.oxml import OxmlElement
    d = docx.Document()
    d.add_paragraph("正文")
    para = d.add_paragraph()
    drawing = OxmlElement('w:drawing')
    para._element.append(drawing)
    d.save(path)


def test_docx_chart_deferred():
    """docx 含图表 drawing (非图片) -> deferred type=chart."""
    in_dir = make_tmp_input({})
    p = os.path.join(in_dir, "chart.docx")
    _make_docx_with_chart(p)
    out = make_tmp_output()
    _, _, e = pipeline.run(in_dir, out)
    assert_eq(e, 0)
    body = open(find_mds(out)[0], encoding="utf-8").read()
    fm = parse_front_matter(body)
    assert_true("deferred" in fm)
    types = [d["type"] for d in fm["deferred"]]
    assert_true("chart" in types, f"应有 chart 类型, got {types}")
    cleanup(in_dir, out)


CASES.append(("docx_chart_deferred", test_docx_chart_deferred))


def test_multiple_deferred_types_one_doc():
    """同一文档有多种 deferred type (formula + chart)."""
    in_dir = make_tmp_input({})
    p = os.path.join(in_dir, "multi.docx")
    import docx
    from docx.oxml import OxmlElement
    d = docx.Document()
    d.add_paragraph("正文")
    # formula
    para1 = d.add_paragraph()
    omath = OxmlElement('m:oMath')
    r = OxmlElement('m:r')
    t = OxmlElement('m:t')
    t.text = "x+1"
    r.append(t)
    omath.append(r)
    para1._element.append(omath)
    # chart
    para2 = d.add_paragraph()
    drawing = OxmlElement('w:drawing')
    para2._element.append(drawing)
    d.save(p)
    out = make_tmp_output()
    _, _, e = pipeline.run(in_dir, out)
    assert_eq(e, 0)
    body = open(find_mds(out)[0], encoding="utf-8").read()
    fm = parse_front_matter(body)
    types = sorted(d["type"] for d in fm["deferred"])
    assert_eq(types, ["chart", "formula"], f"应有 formula+chart, got {types}")
    cleanup(in_dir, out)


CASES.append(("multiple_deferred_types_one_doc", test_multiple_deferred_types_one_doc))


def test_marker_after_image_reference():
    """body 标记紧跟在 ![title](path) 图片引用之后."""
    require_ocr()
    from tests import _samples
    import tempfile
    png_dir = tempfile.mkdtemp()
    png = os.path.join(png_dir, "pic.png")
    _samples.gen_text_png(png, "SOMETEXT")
    try:
        in_dir = make_tmp_input({})
        p = os.path.join(in_dir, "img.docx")
        import docx
        d = docx.Document()
        d.add_paragraph("文字前")
        run = d.add_paragraph().add_run()
        run.add_picture(png)
        d.add_paragraph("文字后")
        d.save(p)
        out = make_tmp_output()
        _, _, e = pipeline.run(in_dir, out)
        assert_eq(e, 0)
        body = open(find_mds(out)[0], encoding="utf-8").read()
        # OCR 成功时: ![](path) 后跟 ``` 代码块
        # OCR 失败时: ![](path) 后跟 <!-- stage2:image-ocr= -->
        assert_true("![" in body, "应有图片引用")
        lines = body.split("\n")
        for i, ln in enumerate(lines):
            if ln.strip().startswith("!["):
                # 下一非空行应是 OCR 文本块或 stage2 标记
                after = "\n".join(lines[i:i+5])
                assert_true("```" in after or "stage2:" in after,
                            "图片引用后应有 OCR 文本或 stage2 标记")
                break
        cleanup(in_dir, out)
    finally:
        cleanup(png_dir)


CASES.append(("marker_after_image_reference", test_marker_after_image_reference))


def test_pure_diagram_image_defers():
    """纯色无文字图片 (OCR 有引擎但返回空) -> defer."""
    require_ocr()
    from PIL import Image
    import tempfile
    png_dir = tempfile.mkdtemp()
    png = os.path.join(png_dir, "solid.png")
    Image.new('RGB', (200, 100), color='blue').save(png)
    try:
        in_dir = make_tmp_input({})
        import shutil
        shutil.copy(png, os.path.join(in_dir, "solid.png"))
        out = make_tmp_output()
        _, _, e = pipeline.run(in_dir, out)
        assert_eq(e, 0)
        body = open(find_mds(out)[0], encoding="utf-8").read()
        fm = parse_front_matter(body)
        assert_true("deferred" in fm, "纯图无文字应 defer")
        types = [d["type"] for d in fm["deferred"]]
        assert_true("image-ocr" in types, f"应 defer image-ocr, got {types}")
        cleanup(in_dir, out)
    finally:
        cleanup(png_dir)


CASES.append(("pure_diagram_image_defers", test_pure_diagram_image_defers))
